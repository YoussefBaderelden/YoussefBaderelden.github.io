"""Build a two-page, single-column Flutter developer CV in PDF and DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parent.parent
PDF_OUT = ROOT / "Youssef_Bader_Flutter_Developer_CV.pdf"
DOCX_OUT = ROOT / "Youssef_Bader_Flutter_Developer_CV.docx"

NAME = "Youssef Hesham Bader Elden"
EMAIL = "yousefhasham512@gmail.com"
PHONE = "+20 114 606 6036"
PORTFOLIO = "https://youssefbaderelden.github.io/"
LINKEDIN = "https://www.linkedin.com/in/yousef-bader-el-den"
GITHUB = "https://github.com/YoussefBaderelden"
MOSTAQL = "https://mostaql.com/u/YoussefBader/portfolio"

pdfmetrics.registerFont(TTFont("Arial", r"C:\Windows\Fonts\arial.ttf"))
pdfmetrics.registerFont(TTFont("Arial-Bold", r"C:\Windows\Fonts\arialbd.ttf"))

SUMMARY = (
    "Flutter developer experienced in building cross-platform mobile applications for Android and iOS, "
    "including production social, streaming, and sports products. Integrates REST APIs, Firebase, live "
    "media services, and payment flows, with hands-on work in mobile UI development and performance "
    "optimization. Contributes from feature implementation through debugging and store releases, with a "
    "focus on reliable, user-centred experiences."
)

EXPERIENCE = [
    {
        "title": "Flutter Developer",
        "company": "TKDS MENA | Remote",
        "dates": "November 2025 - Present",
        "bullets": [
            "Develop Flutter products for a broadcasting and streaming platform, including social live-streaming, mobile broadcasting, and sports media apps released on the App Store and Google Play.",
            "Implemented broadcaster workflows for RTMP, WebRTC, and NDI delivery to OBS and vMix, with camera switching and microphone controls.",
            "Built features for live and on-demand sports, authentication, subscriptions, payments, radio, and real-time updates using Firebase and REST APIs.",
            "Released products across iOS and Android, a Windows video-review app (TKDS VAR), and an Android TV sports app (ASN TV). Also contributed to PlayTooWin (web) and Score Cue."
        ],
    },
    {
        "title": "Flutter Engineer | Contract",
        "company": "BikMedia | iOS and Android",
        "dates": "",
        "bullets": [
            "Integrated Agora RTC, Nosmai effects, a native Game Live fetcher, and a WebRTC filter pipeline for a live social application.",
            "Improved live-room performance through memory-leak cleanup, chat throttling, and RepaintBoundary isolation around video; the app is available on both mobile stores."
        ],
    },
    {
        "title": "Flutter Developer Intern",
        "company": "Digital Egypt Pioneers Initiative (DEPI) | Zagazig, Sharqia",
        "dates": "July 2025 | 6-month program",
        "bullets": ["Studied Java, object-oriented programming, Flutter, and mobile application architecture."],
    },
    {
        "title": "Flutter Developer Intern",
        "company": "Route Academy | Online",
        "dates": "September 2024 | 6-month program",
        "bullets": ["Built Flutter project experience with Provider, Firebase, REST APIs, and Clean Architecture."],
    },
    {
        "title": "Mobile Developer Intern",
        "company": "Digital HUB (D-HUB) | Al-Doqi, Cairo",
        "dates": "June 2024 | 1 month",
        "bullets": ["Studied cross-platform mobile development and user-centred interface design."],
    },
]

SKILLS = [
    ("Programming Languages", "Dart; Java; object-oriented programming"),
    ("Mobile and Desktop Development", "Flutter; Android; iOS; Windows desktop apps; Android TV apps; Flutter web; responsive UI"),
    ("State Management and Architecture", "Provider; Riverpod; GetX; Clean Architecture; BLoC/Cubit (currently learning)"),
    ("Backend and API Integration", "REST APIs; Firebase Authentication; Firestore; Cloud Storage; Firebase Cloud Messaging; Dio; HTTP"),
    ("Databases and Local Storage", "Cloud Firestore; Hive; Shared Preferences"),
    ("Real-Time Communication and Streaming", "Agora RTC; WebRTC; NDI; RTMP; FFmpeg; Socket.IO; camera APIs; audio and video streaming"),
    ("Development Tools and Version Control", "Git; GitHub; Postman; Android Studio; Visual Studio Code; Figma"),
    ("Testing, Deployment, and Performance", "Debugging; performance optimization; App Store, Google Play, and Microsoft Store releases"),
]

PROJECTS = [
    {
        "name": "Fizz Party | TKDS MENA",
        "achievements": [
            "Built social live-streaming features for profiles, live rooms, search, and favourites.",
            "Integrated Firebase and REST-backed app functionality; published the app on iOS and Android.",
        ],
        "tech": "Flutter, Dart, Firebase, REST APIs",
        "links": [("App Store", "https://apps.apple.com/app/id6755402526"), ("Google Play", "https://play.google.com/store/apps/details?id=com.tkds.FizzAPP")],
    },
    {
        "name": "TKDS Streamer | TKDS MENA",
        "achievements": [
            "Developed mobile broadcaster workflows that send phone and iPad feeds to OBS and vMix using RTMP, WebRTC, and NDI.",
            "Implemented camera switching and microphone controls; published the app on iOS and Android.",
        ],
        "tech": "Flutter, Dart, RTMP, WebRTC, NDI, camera APIs",
        "links": [("App Store", "https://apps.apple.com/app/id6755495829"), ("Google Play", "https://play.google.com/store/apps/details?id=com.tkds.tkdsStreamer")],
    },
    {
        "name": "National Golf Network (NGN) | TKDS MENA",
        "achievements": [
            "Implemented live and on-demand golf coverage for high-school, college, and amateur events.",
            "Built premium subscription and payment flows; published the app on iOS and Android.",
        ],
        "tech": "Flutter, Dart, Firebase, video streaming, payments",
        "links": [("App Store", "https://apps.apple.com/app/id6764359434"), ("Google Play", "https://play.google.com/store/apps/details?id=com.tkds.nationalgolfnetwork")],
    },
    {
        "name": "BikMedia",
        "achievements": [
            "Built live-social features with Agora RTC, Nosmai effects, a native Game Live fetcher, and a WebRTC filter pipeline.",
            "Addressed memory leaks, busy-room chat load, and video rendering performance; app is available on iOS and Android.",
        ],
        "tech": "Flutter, Dart, Agora RTC, Nosmai, WebRTC, GetX, Socket.IO",
        "links": [("App Store", "https://apps.apple.com/app/id6756263985"), ("Google Play", "https://play.google.com/store/apps/details?id=com.bikmedianewapp.vr")],
    },
    {
        "name": "TKDS VAR | TKDS MENA",
        "achievements": [
            "Contributed NDI and FFmpeg-based video review for live-sports camera feeds on Windows.",
            "Implemented multi-camera comparison, timeline navigation, and timestamped tags; published on the Microsoft Store.",
        ],
        "tech": "Flutter, Dart, FFmpeg, media_kit, NDI",
        "links": [("Microsoft Store", "https://apps.microsoft.com/detail/9nhp3hphp9t3")],
    },
    {
        "name": "PlayTooWin | TKDS MENA",
        "achievements": [
            "Contributed to a Flutter web live-bingo product with sign-in and card-purchasing flows.",
            "Implemented the player experience for following the live stream.",
        ],
        "tech": "Flutter, Dart, Riverpod, live video",
        "links": [("Website", "https://playtoowin.com/")],
    },
    {
        "name": "Pasty",
        "achievements": [
            "Built a cake-customization and sweets-commerce app with customer and designer experiences, Arabic/English support, maps, and delivery QR flows.",
            "Released the app on iOS and Android with local payment options.",
        ],
        "tech": "Flutter, Dart, BLoC, Hive, maps, payments",
        "links": [("App Store", "https://apps.apple.com/us/app/basty-باستي/id6766212636?platform=mac"), ("Google Play", "https://play.google.com/store/apps/details?id=com.basty.app")],
    },
]

MORE_PROJECTS = [
    ("TKDS Sports", "https://apps.apple.com/app/id6778073209"),
    ("TKDS Sports on Google Play", "https://play.google.com/store/apps/details?id=com.tkds.tkdssports"),
    ("TKDS Radio", "https://play.google.com/store/apps/details?id=com.tkds.RadioAPP"),
    ("Maadi Sports", "https://play.google.com/store/apps/details?id=com.tkds.maadisports"),
    ("ASN TV (Android TV)", "https://play.google.com/store/apps/details?id=com.tkds.antlersn"),
    ("Sports Flix", "https://play.google.com/store/apps/details?id=com.tkds.sportsfixlive"),
    ("Southwest Ohio", "https://play.google.com/store/apps/details?id=com.tkds.southwestohio"),
    ("Score Cue", ""),
    ("News", "https://github.com/YoussefBaderelden/news"),
    ("Evently", "https://github.com/YoussefBaderelden/evently"),
    ("Movie Browser", "https://github.com/NadaAymanElsayed/movies"),
    ("Islamic App", "https://github.com/YoussefBaderelden/islamic_app"),
    ("Fitness App", "https://github.com/YoussefBaderelden/fitness_app"),
]


def pdf_style(name, **overrides):
    values = dict(fontName="Arial", fontSize=10, leading=12.6, textColor="#20252B", alignment=TA_LEFT)
    values.update(overrides)
    return ParagraphStyle(name, **values)


PDF_NAME = pdf_style("name", fontName="Arial-Bold", fontSize=22, leading=26, textColor="#163D4B", spaceAfter=1)
PDF_ROLE = pdf_style("role", fontSize=12, leading=15, textColor="#36535B", spaceAfter=5)
PDF_CONTACT = pdf_style("contact", fontSize=9.3, leading=11.5, spaceAfter=2)
PDF_HEAD = pdf_style("head", fontName="Arial-Bold", fontSize=12.5, leading=15, textColor="#163D4B", spaceBefore=8, spaceAfter=3)
PDF_JOB = pdf_style("job", fontName="Arial-Bold", fontSize=10.2, leading=13, spaceBefore=4, spaceAfter=1)
PDF_META = pdf_style("meta", fontSize=9.2, leading=11.5, textColor="#53646A", spaceAfter=1)
PDF_BODY = pdf_style("body", fontSize=10.2, leading=13.1, spaceAfter=2)
PDF_BULLET = pdf_style("bullet", fontSize=9.8, leading=12.6, leftIndent=10, firstLineIndent=-7, spaceAfter=1)
PDF_PROJECT = pdf_style("project", fontName="Arial-Bold", fontSize=10.2, leading=13, spaceBefore=4, spaceAfter=1)
PDF_SMALL = pdf_style("small", fontSize=9.5, leading=12, spaceAfter=1)


def pdf_link(url, label):
    return f'<link href="{url}" color="#147E83"><u>{label}</u></link>'


def pdf_heading(text):
    return KeepTogether([Paragraph(text, PDF_HEAD), HRFlowable(width="100%", thickness=0.5, color="#77A6A2", spaceBefore=1, spaceAfter=3)])


def pdf_bullet(text):
    return Paragraph(f"- {text}", PDF_BULLET)


def pdf_story():
    story = [
        Paragraph(NAME, PDF_NAME),
        Paragraph("Flutter Developer | Cross-Platform Mobile Applications", PDF_ROLE),
        Paragraph(f"Cairo, Egypt | {PHONE} | {pdf_link('mailto:' + EMAIL, EMAIL)}", PDF_CONTACT),
        Paragraph("Portfolio: " + pdf_link(PORTFOLIO, "youssefbaderelden.github.io") + " | LinkedIn: " + pdf_link(LINKEDIN, "Profile") + " | GitHub: " + pdf_link(GITHUB, "Profile") + " | Mostaql: " + pdf_link(MOSTAQL, "Portfolio"), PDF_CONTACT),
        Paragraph("Arabic: Native | English: Intermediate", PDF_CONTACT),
        pdf_heading("Professional Summary"),
        Paragraph(SUMMARY, PDF_BODY),
        pdf_heading("Professional Experience"),
    ]
    for item in EXPERIENCE:
        meta = item["company"] + (" | " + item["dates"] if item["dates"] else "")
        story.extend([Paragraph(item["title"], PDF_JOB), Paragraph(meta, PDF_META)])
        story.extend(pdf_bullet(point) for point in item["bullets"])
    story.extend([
        pdf_heading("Education"),
        Paragraph("Engineering studies (in progress)", PDF_JOB),
        Paragraph("Faculty of Engineering, Zagazig University | Engineering Department", PDF_META),
        Paragraph("October 2022 - August 2027 (expected) | Third-year student | Very Good", PDF_BODY),
        PageBreak(),
        pdf_heading("Selected Projects"),
    ])
    for item in PROJECTS:
        story.append(Paragraph(item["name"], PDF_PROJECT))
        story.extend(pdf_bullet(point) for point in item["achievements"])
        links = " | ".join(pdf_link(url, label) for label, url in item["links"])
        story.append(Paragraph(f"<b>Technologies:</b> {item['tech']} | <b>Links:</b> {links}", PDF_SMALL))
    more_links = " | ".join(pdf_link(url, label) if url else label for label, url in MORE_PROJECTS)
    story.append(pdf_heading("More Projects"))
    story.append(Paragraph(more_links, PDF_SMALL))
    story.append(pdf_heading("Technical Skills"))
    for category, skills in SKILLS:
        story.append(Paragraph(f"<b>{category}:</b> {skills}", PDF_SMALL))
    return story


def add_hyperlink(paragraph, label, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "147E83")
    properties.append(color)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_docx_style(document):
    section = document.sections[0]
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.52)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(32, 37, 43)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0


def docx_heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11.5)
    run.font.color.rgb = RGBColor(22, 61, 75)
    return paragraph


def docx_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.add_run(text)
    return paragraph


def docx_story():
    document = Document()
    set_docx_style(document)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(NAME)
    run.bold = True
    run.font.size = Pt(19)
    run.font.color.rgb = RGBColor(22, 61, 75)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run("Flutter Developer | Cross-Platform Mobile Applications").font.size = Pt(10)
    document.add_paragraph(f"Cairo, Egypt | {PHONE} | {EMAIL}")
    paragraph = document.add_paragraph()
    for index, (label, url) in enumerate((("Portfolio", PORTFOLIO), ("LinkedIn", LINKEDIN), ("GitHub", GITHUB), ("Mostaql", MOSTAQL))):
        if index:
            paragraph.add_run(" | ")
        add_hyperlink(paragraph, label, url)
    document.add_paragraph("Arabic: Native | English: Intermediate")
    docx_heading(document, "Professional Summary")
    document.add_paragraph(SUMMARY)
    docx_heading(document, "Professional Experience")
    for item in EXPERIENCE:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(item["title"])
        run.bold = True
        meta = item["company"] + (" | " + item["dates"] if item["dates"] else "")
        paragraph = document.add_paragraph(meta)
        paragraph.paragraph_format.space_after = Pt(1)
        for point in item["bullets"]:
            docx_bullet(document, point)
    docx_heading(document, "Education")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run("Engineering studies (in progress)")
    run.bold = True
    document.add_paragraph("Faculty of Engineering, Zagazig University | Engineering Department")
    document.add_paragraph("October 2022 - August 2027 (expected) | Third-year student | Very Good")
    document.add_page_break()
    docx_heading(document, "Selected Projects")
    for item in PROJECTS:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(item["name"])
        run.bold = True
        for point in item["achievements"]:
            docx_bullet(document, point)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run("Technologies: ")
        run.bold = True
        paragraph.add_run(item["tech"] + " | Links: ")
        for index, (label, url) in enumerate(item["links"]):
            if index:
                paragraph.add_run(" | ")
            add_hyperlink(paragraph, label, url)
    docx_heading(document, "More Projects")
    paragraph = document.add_paragraph()
    for index, (label, url) in enumerate(MORE_PROJECTS):
        if index:
            paragraph.add_run(" | ")
        if url:
            add_hyperlink(paragraph, label, url)
        else:
            paragraph.add_run(label)
    docx_heading(document, "Technical Skills")
    for category, skills in SKILLS:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(category + ": ")
        run.bold = True
        paragraph.add_run(skills)
    document.core_properties.title = "Youssef Bader - Flutter Developer CV"
    document.core_properties.author = NAME
    document.save(DOCX_OUT)


def main():
    pdf = SimpleDocTemplate(
        str(PDF_OUT), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=11 * mm, bottomMargin=11 * mm,
        title="Youssef Bader - Flutter Developer CV", author=NAME,
    )
    pdf.build(pdf_story())
    docx_story()
    print(PDF_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    main()
